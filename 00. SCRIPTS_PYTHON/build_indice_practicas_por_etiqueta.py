#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_indice_practicas_por_etiqueta.py
Genera el índice exhaustivo de todas las prácticas de las ternas agrupadas por etiqueta (categoría)
en tres formatos:
1. INDICE_PRACTICAS_POR_ETIQUETA.html (Visual, interactivo, con buscador en vivo y filtros)
2. INDICE_PRACTICAS_POR_ETIQUETA.docx (Documento Word maquetado para imprimir y leer)
3. INDICE_PRACTICAS_POR_ETIQUETA.md   (Formato Markdown para GitHub y documentación)

Ubicación de salida:
CURSO-IA/INDICE_DE_PRACTICAS_POR_ETIQUETA/
"""

import os
import csv
import re
import urllib.parse
import unicodedata
import json
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

ROOT_DIR = "/Users/externo/Library/Mobile Documents/com~apple~CloudDocs/PERSONAL/CLASES DE TECNOLOGÍA/CURSO-IA"
CSV_PATH = os.path.join(ROOT_DIR, "PANELES_CSV", "IA_INDICE_COMPLETO_60_SESIONES.csv")
OUTPUT_DIR = os.path.join(ROOT_DIR, "INDICE_DE_PRACTICAS_POR_ETIQUETA")
os.makedirs(OUTPUT_DIR, exist_ok=True)

HTML_OUT = os.path.join(OUTPUT_DIR, "INDICE_PRACTICAS_POR_ETIQUETA.html")
DOCX_OUT = os.path.join(OUTPUT_DIR, "INDICE_PRACTICAS_POR_ETIQUETA.docx")
MD_OUT   = os.path.join(OUTPUT_DIR, "INDICE_PRACTICAS_POR_ETIQUETA.md")

TAG_COLORS = {
    "TXT":   {"bg": "#1e3a8a", "text": "#93c5fd", "border": "#3b82f6", "name": "Prompts de Comunicación y Texto"},
    "EST":   {"bg": "#581c87", "text": "#d8b4fe", "border": "#a855f7", "name": "Estilos Visuales de Imagen"},
    "PRAC":  {"bg": "#064e3b", "text": "#6ee7b7", "border": "#10b981", "name": "Talleres Prácticos con Gemini"},
    "FRAC":  {"bg": "#701a75", "text": "#f472b6", "border": "#ec4899", "name": "Fractales en IA (Vídeos y Fichas)"},
    "INT":   {"bg": "#7c2d12", "text": "#fdba74", "border": "#f97316", "name": "El Mundo por Dentro y Reconstrucción"},
    "FUT":   {"bg": "#134e4a", "text": "#5eead4", "border": "#14b8a6", "name": "Línea de Tiempo del Futuro (2030+)"},
    "NAT":   {"bg": "#14532d", "text": "#86efac", "border": "#22c55e", "name": "Naturaleza Fascinante y Biodiversidad"},
    "ARTE":  {"bg": "#831843", "text": "#f9a8d4", "border": "#f43f5e", "name": "Obras Maestras del Arte Universal"},
    "NIV":   {"bg": "#312e81", "text": "#a5b4fc", "border": "#6366f1", "name": "Escalafones y Niveles (Cultura 101)"},
    "TRUC":  {"bg": "#713f12", "text": "#fde047", "border": "#eab308", "name": "Trucos y Soluciones Cotidianas"},
    "CUENT": {"bg": "#881337", "text": "#fda4af", "border": "#f43f5e", "name": "Cuentos Ilustrados para Nietos"},
    "MOVIL": {"bg": "#1e293b", "text": "#38bdf8", "border": "#0ea5e9", "name": "El Salvavidas del Móvil (Cámara y Voz)"},
    "MEM":   {"bg": "#451a03", "text": "#fdba74", "border": "#d97706", "name": "Cápsula de la Memoria"},
    "MEC":   {"bg": "#3f3f46", "text": "#e4e4e7", "border": "#a1a1aa", "name": "Cómo Funcionan las Cosas (Vídeo 3D)"},
    "GUIA":  {"bg": "#18181b", "text": "#fafafa", "border": "#71717a", "name": "Glosario de Etiquetas y Guía"},
}

def parse_csv():
    categories = {}
    with open(CSV_PATH, 'r', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        for row in reader:
            tag_theme = row['TEMA_CLASSROOM'].strip()
            title = row['TITULO_MATERIAL'].strip()
            desc = row['DESCRIPCION_MATERIAL'].strip()
            url = row['URL_GITHUB'].strip()
            
            # Extract Tag Code e.g. "TXT" from "01. [TXT] Prompts de Comunicación y Texto"
            match_code = re.search(r'\[(.*?)\]', tag_theme)
            tag_code = match_code.group(1) if match_code else "OTROS"
            
            # Extract Session number e.g. "[Sesión 01]" or from title
            match_session = re.search(r'\[Sesión\s*(\d+)\]', title, re.IGNORECASE)
            session_num = match_session.group(1) if match_session else "00"
            
            # Clean title
            clean_t = re.sub(r'\[Sesión\s*\d+\]\s*', '', title).strip()
            
            # Type of file (PDF or MP4)
            file_type = "MP4" if "(MP4)" in title.upper() or url.lower().endswith(".mp4") else "PDF"
            
            # Local relative link from INDICE_DE_PRACTICAS_POR_ETIQUETA/
            # URL_GITHUB format: https://javiercursocorreo-gif.github.io/Curso-IA/CLASES/...
            rel_path = ""
            if "Curso-IA/" in url:
                sub_rel = urllib.parse.unquote(url.split("Curso-IA/")[1])
                rel_path = os.path.join("..", sub_rel)
                
            item = {
                "session": session_num,
                "title_raw": title,
                "title_clean": clean_t,
                "desc": desc,
                "url": url,
                "rel_path": rel_path,
                "file_type": file_type,
                "tag_code": tag_code
            }
            
            if tag_theme not in categories:
                categories[tag_theme] = {
                    "tag_code": tag_code,
                    "theme_name": tag_theme,
                    "items": []
                }
            categories[tag_theme]["items"].append(item)
            
    return categories

def build_markdown(categories):
    md = []
    md.append("# 📚 ÍNDICE MAESTRO DE PRÁCTICAS Y TERNAS POR ETIQUETA")
    md.append("## Curso de Inteligencia Artificial para Mayores de 60 Años\n")
    md.append("> **Catálogo General**: 14 Categorías Temáticas • 60 Sesiones • 1.059 Prácticas y Recursos Didácticos.\n")
    md.append("---\n")
    
    # Tabla resumen de categorías
    md.append("### 📊 Resumen por Categorías / Etiquetas\n")
    md.append("| Etiqueta | Nombre del Área | N.º Prácticas |")
    md.append("| :--- | :--- | :---: |")
    total_items = 0
    for theme, data in categories.items():
        count = len(data["items"])
        total_items += count
        md.append(f"| **`[{data['tag_code']}]`** | {theme} | **{count}** |")
    md.append(f"| **TOTAL** | **Todas las categorías** | **{total_items}** |\n")
    md.append("---\n")
    
    # Detalle por categoría
    for theme, data in categories.items():
        code = data["tag_code"]
        items = data["items"]
        md.append(f"## 🏷️ {theme} ({len(items)} prácticas)\n")
        md.append("| Sesión | Título de la Práctica / Ficha | Tipo | Enlace Directo |")
        md.append("| :---: | :--- | :---: | :---: |")
        for it in items:
            ses_str = f"Sesión {it['session']}" if it['session'] != "00" else "Guía"
            md.append(f"| `{ses_str}` | {it['title_clean']} | **{it['file_type']}** | [Abrir Archivo]({it['url']}) |")
        md.append("\n---\n")
        
    with open(MD_OUT, 'w', encoding='utf-8') as f:
        f.write("\n".join(md))
    print(f"✅ Markdown generado en: {MD_OUT}")

def build_html(categories):
    total_items = sum(len(c["items"]) for c in categories.values())
    
    # Generate category navigation buttons
    nav_buttons = []
    for theme, data in categories.items():
        code = data["tag_code"]
        c_info = TAG_COLORS.get(code, {"bg": "#334155", "text": "#f8fafc", "border": "#64748b"})
        count = len(data["items"])
        nav_buttons.append(f'''
            <button class="filter-btn" data-tag="{code}" onclick="filterTag('{code}')" style="--btn-border: {c_info['border']}; --btn-text: {c_info['text']};">
                <span class="btn-tag">[{code}]</span>
                <span class="btn-name">{code}</span>
                <span class="btn-count">{count}</span>
            </button>
        ''')
        
    # Generate items markup
    sections_html = []
    for theme, data in categories.items():
        code = data["tag_code"]
        c_info = TAG_COLORS.get(code, {"bg": "#334155", "text": "#94a3b8", "border": "#64748b"})
        items = data["items"]
        
        cards_html = []
        for it in items:
            badge_type_class = "badge-mp4" if it['file_type'] == "MP4" else "badge-pdf"
            ses_badge = f"Sesión {it['session']}" if it['session'] != "00" else "Guía Oficial"
            cards_html.append(f'''
                <div class="practice-card" data-tag="{code}" data-title="{it['title_clean'].lower()}" data-session="{it['session']}">
                    <div class="card-meta">
                        <span class="tag-pill" style="background: {c_info['bg']}; color: {c_info['text']}; border-color: {c_info['border']};">[{code}]</span>
                        <span class="session-pill">{ses_badge}</span>
                        <span class="type-pill {badge_type_class}">{it['file_type']}</span>
                    </div>
                    <h3 class="card-title">{it['title_clean']}</h3>
                    <p class="card-desc">{it['desc']}</p>
                    <div class="card-actions">
                        <a href="{it['rel_path']}" target="_blank" class="btn-action btn-local" title="Abrir archivo localmente en este ordenador">📁 Abrir Local</a>
                        <a href="{it['url']}" target="_blank" class="btn-action btn-cloud" title="Ver en GitHub Pages (Online)">☁️ Ver Online</a>
                    </div>
                </div>
            ''')
            
        sections_html.append(f'''
            <section class="tag-section" id="sec-{code}" data-tag="{code}">
                <div class="section-header" style="border-left-color: {c_info['border']};">
                    <div class="header-left">
                        <span class="section-tag" style="color: {c_info['text']};">[{code}]</span>
                        <h2 class="section-title">{theme}</h2>
                    </div>
                    <span class="section-counter">{len(items)} prácticas</span>
                </div>
                <div class="cards-grid">
                    {''.join(cards_html)}
                </div>
            </section>
        ''')

    html_content = f'''<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Índice de Prácticas por Etiqueta • Curso de IA</title>
    <style>
        :root {{
            --bg-base: #0a0e17;
            --bg-surface: #111827;
            --bg-card: rgba(17, 24, 39, 0.85);
            --border-subtle: rgba(255, 255, 255, 0.08);
            --border-focus: #38bdf8;
            --text-main: #f9fafb;
            --text-sub: #9ca3af;
            --text-muted: #6b7280;
            --primary: #38bdf8;
        }}
        * {{ box-sizing: border-box; margin: 0; padding: 0; }}
        body {{
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
            background: var(--bg-base);
            color: var(--text-main);
            line-height: 1.5;
            padding-bottom: 80px;
        }}
        
        /* HEADER */
        .hero {{
            background: radial-gradient(circle at 50% 0%, #1e293b 0%, #0a0e17 80%);
            padding: 48px 24px 32px;
            text-align: center;
            border-bottom: 1px solid var(--border-subtle);
        }}
        .hero-badge {{
            display: inline-block;
            background: rgba(56, 189, 248, 0.12);
            color: #38bdf8;
            border: 1px solid rgba(56, 189, 248, 0.3);
            border-radius: 9999px;
            padding: 4px 16px;
            font-size: 0.85rem;
            font-weight: 600;
            margin-bottom: 16px;
            letter-spacing: 0.5px;
        }}
        .hero-title {{
            font-size: 2.2rem;
            font-weight: 800;
            background: linear-gradient(135deg, #ffffff 40%, #94a3b8 100%);
            -webkit-background-clip: text;
            -webkit-text-fill-color: transparent;
            margin-bottom: 10px;
        }}
        .hero-subtitle {{
            font-size: 1.05rem;
            color: var(--text-sub);
            max-width: 800px;
            margin: 0 auto 24px;
        }}
        .stats-bar {{
            display: flex;
            justify-content: center;
            gap: 24px;
            flex-wrap: wrap;
            margin-top: 16px;
        }}
        .stat-item {{
            background: var(--bg-surface);
            border: 1px solid var(--border-subtle);
            border-radius: 12px;
            padding: 10px 20px;
            text-align: center;
        }}
        .stat-val {{
            font-size: 1.3rem;
            font-weight: 700;
            color: #38bdf8;
            display: block;
        }}
        .stat-lbl {{
            font-size: 0.78rem;
            color: var(--text-muted);
            text-transform: uppercase;
            letter-spacing: 0.5px;
        }}

        /* STICKY TOOLBAR */
        .toolbar {{
            position: sticky;
            top: 0;
            z-index: 100;
            background: rgba(10, 14, 23, 0.95);
            backdrop-filter: blur(12px);
            border-bottom: 1px solid var(--border-subtle);
            padding: 14px 24px;
            display: flex;
            flex-direction: column;
            gap: 12px;
            max-width: 1400px;
            margin: 0 auto;
        }}
        .search-row {{
            display: flex;
            gap: 12px;
            align-items: center;
        }}
        .search-box {{
            flex: 1;
            position: relative;
        }}
        .search-input {{
            width: 100%;
            background: var(--bg-surface);
            border: 1px solid var(--border-subtle);
            border-radius: 10px;
            padding: 12px 16px 12px 42px;
            font-size: 0.95rem;
            color: #fff;
            outline: none;
            transition: all 0.2s;
        }}
        .search-input:focus {{
            border-color: var(--border-focus);
            box-shadow: 0 0 0 3px rgba(56, 189, 248, 0.2);
        }}
        .search-icon {{
            position: absolute;
            left: 14px;
            top: 50%;
            transform: translateY(-50%);
            font-size: 1.1rem;
            color: var(--text-muted);
        }}
        .clear-btn {{
            background: #374151;
            border: none;
            color: #d1d5db;
            padding: 10px 18px;
            border-radius: 8px;
            cursor: pointer;
            font-weight: 600;
            font-size: 0.85rem;
        }}
        .clear-btn:hover {{ background: #4b5563; }}

        /* FILTER BUTTONS */
        .filters-scroll {{
            display: flex;
            gap: 8px;
            overflow-x: auto;
            padding-bottom: 4px;
            scrollbar-width: thin;
        }}
        .filter-btn {{
            display: inline-flex;
            align-items: center;
            gap: 6px;
            background: var(--bg-surface);
            border: 1px solid var(--border-subtle);
            color: var(--text-sub);
            padding: 6px 12px;
            border-radius: 8px;
            font-size: 0.82rem;
            font-weight: 600;
            cursor: pointer;
            white-space: nowrap;
            transition: all 0.18s;
        }}
        .filter-btn:hover {{
            border-color: var(--btn-border);
            color: var(--btn-text);
            transform: translateY(-1px);
        }}
        .filter-btn.active {{
            background: rgba(56, 189, 248, 0.15);
            border-color: #38bdf8;
            color: #38bdf8;
        }}
        .filter-btn .btn-count {{
            background: rgba(255, 255, 255, 0.08);
            padding: 2px 6px;
            border-radius: 6px;
            font-size: 0.75rem;
        }}

        /* CONTAINER */
        .main-container {{
            max-width: 1400px;
            margin: 32px auto 0;
            padding: 0 24px;
        }}

        /* SECTIONS */
        .tag-section {{
            margin-bottom: 44px;
            scroll-margin-top: 130px;
        }}
        .section-header {{
            display: flex;
            justify-content: space-between;
            align-items: center;
            padding: 10px 16px;
            background: var(--bg-surface);
            border-radius: 10px;
            border-left: 5px solid #38bdf8;
            margin-bottom: 20px;
        }}
        .header-left {{
            display: flex;
            align-items: center;
            gap: 10px;
        }}
        .section-tag {{
            font-weight: 800;
            font-size: 1.1rem;
            font-family: monospace;
        }}
        .section-title {{
            font-size: 1.15rem;
            font-weight: 700;
            color: #fff;
        }}
        .section-counter {{
            background: rgba(255, 255, 255, 0.06);
            color: var(--text-sub);
            padding: 4px 10px;
            border-radius: 8px;
            font-size: 0.82rem;
            font-weight: 600;
        }}

        /* CARDS GRID */
        .cards-grid {{
            display: grid;
            grid-template-columns: repeat(auto-fill, minmax(320px, 1fr));
            gap: 16px;
        }}
        .practice-card {{
            background: var(--bg-card);
            border: 1px solid var(--border-subtle);
            border-radius: 12px;
            padding: 16px;
            display: flex;
            flex-direction: column;
            transition: all 0.2s ease;
        }}
        .practice-card:hover {{
            border-color: rgba(56, 189, 248, 0.4);
            transform: translateY(-2px);
            box-shadow: 0 8px 24px rgba(0, 0, 0, 0.4);
        }}
        .card-meta {{
            display: flex;
            gap: 8px;
            align-items: center;
            margin-bottom: 10px;
        }}
        .tag-pill {{
            font-size: 0.75rem;
            font-weight: 700;
            padding: 2px 8px;
            border-radius: 6px;
            border: 1px solid;
            font-family: monospace;
        }}
        .session-pill {{
            background: #1f2937;
            color: #9ca3af;
            font-size: 0.75rem;
            font-weight: 600;
            padding: 2px 8px;
            border-radius: 6px;
        }}
        .type-pill {{
            font-size: 0.72rem;
            font-weight: 700;
            padding: 2px 6px;
            border-radius: 4px;
            margin-left: auto;
        }}
        .badge-pdf {{ background: #dc2626; color: #fee2e2; }}
        .badge-mp4 {{ background: #7c3aed; color: #ede9fe; }}

        .card-title {{
            font-size: 0.96rem;
            font-weight: 600;
            color: #f3f4f6;
            margin-bottom: 8px;
            line-height: 1.35;
        }}
        .card-desc {{
            font-size: 0.82rem;
            color: var(--text-muted);
            margin-bottom: 14px;
            flex: 1;
            display: -webkit-box;
            -webkit-line-clamp: 2;
            -webkit-box-orient: vertical;
            overflow: hidden;
        }}
        .card-actions {{
            display: flex;
            gap: 8px;
            margin-top: auto;
        }}
        .btn-action {{
            flex: 1;
            text-align: center;
            padding: 8px 10px;
            border-radius: 8px;
            font-size: 0.8rem;
            font-weight: 600;
            text-decoration: none;
            transition: all 0.15s;
        }}
        .btn-local {{
            background: #1e293b;
            color: #38bdf8;
            border: 1px solid rgba(56, 189, 248, 0.3);
        }}
        .btn-local:hover {{
            background: #0284c7;
            color: #fff;
        }}
        .btn-cloud {{
            background: rgba(255, 255, 255, 0.05);
            color: #94a3b8;
            border: 1px solid var(--border-subtle);
        }}
        .btn-cloud:hover {{
            background: rgba(255, 255, 255, 0.12);
            color: #fff;
        }}

        .empty-results {{
            display: none;
            text-align: center;
            padding: 60px 20px;
            color: var(--text-muted);
            font-size: 1.1rem;
        }}
    </style>
</head>
<body>

    <header class="hero">
        <span class="hero-badge">CURSO DE INTELIGENCIA ARTIFICIAL</span>
        <h1 class="hero-title">Índice Maestro de Prácticas por Etiqueta</h1>
        <p class="hero-subtitle">
            Catálogo completo e interactivo de todas las ternas didácticas organizadas por su bloque pedagógico oficial.
            Puedes buscar por título, número de sesión o filtrar directamente por etiquetas.
        </p>
        <div class="stats-bar">
            <div class="stat-item">
                <span class="stat-val">{len(categories)}</span>
                <span class="stat-lbl">Etiquetas Oficiales</span>
            </div>
            <div class="stat-item">
                <span class="stat-val">60</span>
                <span class="stat-lbl">Sesiones del Curso</span>
            </div>
            <div class="stat-item">
                <span class="stat-val">{total_items}</span>
                <span class="stat-lbl">Prácticas y Fichas</span>
            </div>
        </div>
    </header>

    <div class="toolbar">
        <div class="search-row">
            <div class="search-box">
                <span class="search-icon">🔍</span>
                <input type="text" id="searchInput" class="search-input" placeholder="Buscar práctica por nombre, palabra clave o sesión (ej: faro, gato, sesión 15, fractal, 008)..." oninput="handleSearch()">
            </div>
            <button class="clear-btn" onclick="resetFilters()">Mostrar Todo</button>
        </div>
        <div class="filters-scroll">
            <button class="filter-btn active" data-tag="ALL" onclick="filterTag('ALL')">
                <span>⭐ Todas las Etiquetas</span>
                <span class="btn-count">{total_items}</span>
            </button>
            {''.join(nav_buttons)}
        </div>
    </div>

    <main class="main-container" id="mainContainer">
        {''.join(sections_html)}
        <div class="empty-results" id="emptyResults">
            <p>😕 No se han encontrado prácticas que coincidan con la búsqueda.</p>
        </div>
    </main>

    <script>
        let currentTag = 'ALL';

        function filterTag(tag) {{
            currentTag = tag;
            document.querySelectorAll('.filter-btn').forEach(b => {{
                if (b.dataset.tag === tag) {{
                    b.classList.add('active');
                }} else {{
                    b.classList.remove('active');
                }}
            }});
            
            if (tag !== 'ALL') {{
                const targetSec = document.getElementById('sec-' + tag);
                if (targetSec) {{
                    targetSec.scrollIntoView({{ behavior: 'smooth' }});
                }}
            }}
            applyFilters();
        }}

        function handleSearch() {{
            applyFilters();
        }}

        function resetFilters() {{
            document.getElementById('searchInput').value = '';
            filterTag('ALL');
            window.scrollTo({{ top: 0, behavior: 'smooth' }});
        }}

        function applyFilters() {{
            const query = document.getElementById('searchInput').value.trim().toLowerCase();
            let totalVisible = 0;

            document.querySelectorAll('.tag-section').forEach(sec => {{
                const secTag = sec.dataset.tag;
                const matchTag = (currentTag === 'ALL' || currentTag === secTag);
                
                let secVisibleCards = 0;
                sec.querySelectorAll('.practice-card').forEach(card => {{
                    const title = card.dataset.title;
                    const session = card.dataset.session;
                    const tag = card.dataset.tag.toLowerCase();
                    const fullText = title + " sesion " + session + " " + tag;
                    
                    const matchQuery = !query || fullText.includes(query);
                    
                    if (matchTag && matchQuery) {{
                        card.style.display = 'flex';
                        secVisibleCards++;
                        totalVisible++;
                    }} else {{
                        card.style.display = 'none';
                    }}
                }});

                if (secVisibleCards > 0) {{
                    sec.style.display = 'block';
                }} else {{
                    sec.style.display = 'none';
                }}
            }});

            const emptyEl = document.getElementById('emptyResults');
            if (totalVisible === 0) {{
                emptyEl.style.display = 'block';
            }} else {{
                emptyEl.style.display = 'none';
            }}
        }}
    </script>
</body>
</html>
'''
    with open(HTML_OUT, 'w', encoding='utf-8') as f:
        f.write(html_content)
    print(f"✅ HTML interactivo generado en: {HTML_OUT}")

def build_docx(categories):
    doc = docx.Document()
    
    # Page Margins
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        
    styles = doc.styles
    
    # Title
    p_title = doc.add_paragraph('ÍNDICE MAESTRO DE PRÁCTICAS POR ETIQUETA')
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.runs[0]
    r_t.font.name = 'Calibri'
    r_t.font.size = Pt(22)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    
    # Subtitle
    p_sub = doc.add_paragraph('Curso de Inteligencia Artificial para Mayores de 60 Años • Catálogo de Ternas Prácticas')
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.runs[0]
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(12)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    
    doc.add_paragraph()
    
    # Intro box
    p_intro = doc.add_paragraph(
        'Este documento recoge la totalidad de las fichas, prácticas y recursos multimedia creados para las 60 sesiones '
        'del curso, organizados metódicamente por sus 14 Etiquetas Temáticas Oficiales para facilitar la consulta docente y la preparación de clases.'
    )
    p_intro.paragraph_format.space_after = Pt(14)
    
    # Summary Table
    total_items = sum(len(c["items"]) for c in categories.values())
    t_summary = doc.add_table(rows=1, cols=3)
    t_summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = t_summary.rows[0].cells
    hdr_cells[0].text = 'CÓDIGO'
    hdr_cells[1].text = 'ÁREA / CATEGORÍA TEMÁTICA'
    hdr_cells[2].text = 'N.º PRÁCTICAS'
    for c in hdr_cells:
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shd = parse_xml(r'<w:shd {} w:fill="0F172A"/>'.format(nsdecls('w')))
        c._tc.get_or_add_tcPr().append(shd)
        
    for theme, data in categories.items():
        row_cells = t_summary.add_row().cells
        row_cells[0].text = f"[{data['tag_code']}]"
        row_cells[1].text = theme
        row_cells[2].text = str(len(data['items']))
        row_cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        
    # Total row
    tot_row = t_summary.add_row().cells
    tot_row[0].text = "TOTAL"
    tot_row[1].text = "Todas las 14 Etiquetas y Guía Oficial"
    tot_row[2].text = str(total_items)
    for c in tot_row:
        c.paragraphs[0].runs[0].font.bold = True
        shd = parse_xml(r'<w:shd {} w:fill="E2E8F0"/>'.format(nsdecls('w')))
        c._tc.get_or_add_tcPr().append(shd)
        
    doc.add_page_break()
    
    # Detailed Sections by Tag
    for theme, data in categories.items():
        code = data["tag_code"]
        items = data["items"]
        
        # Heading 1
        h1 = doc.add_heading(level=1)
        r_h1 = h1.add_run(f"Etiqueta [{code}] — {theme} ({len(items)} prácticas)")
        r_h1.font.name = 'Calibri'
        r_h1.font.size = Pt(14)
        r_h1.font.bold = True
        r_h1.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        
        # Table of items
        t_items = doc.add_table(rows=1, cols=4)
        t_items.alignment = WD_TABLE_ALIGNMENT.CENTER
        th_cells = t_items.rows[0].cells
        th_cells[0].text = 'Sesión'
        th_cells[1].text = 'Título de la Práctica / Ficha'
        th_cells[2].text = 'Tipo'
        th_cells[3].text = 'Enlace Web'
        for cell in th_cells:
            cell.paragraphs[0].runs[0].font.bold = True
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            shd = parse_xml(r'<w:shd {} w:fill="1E293B"/>'.format(nsdecls('w')))
            cell._tc.get_or_add_tcPr().append(shd)
            
        for it in items:
            row = t_items.add_row().cells
            row[0].text = f"S.{it['session']}" if it['session'] != "00" else "Guía"
            row[1].text = it['title_clean']
            row[2].text = it['file_type']
            row[3].text = it['url']
            row[0].paragraphs[0].runs[0].font.size = Pt(9)
            row[1].paragraphs[0].runs[0].font.size = Pt(9)
            row[2].paragraphs[0].runs[0].font.size = Pt(9)
            row[3].paragraphs[0].runs[0].font.size = Pt(8)
            
        doc.add_paragraph() # spacing
        
    doc.save(DOCX_OUT)
    print(f"✅ Word (DOCX) generado en: {DOCX_OUT}")

def main():
    print("🚀 Iniciando generación del Índice de Prácticas por Etiqueta...")
    categories = parse_csv()
    build_markdown(categories)
    build_html(categories)
    build_docx(categories)
    print("🎉 ¡Proceso completado con éxito!")

if __name__ == "__main__":
    main()
