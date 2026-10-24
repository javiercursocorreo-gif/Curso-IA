#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_guia_docx_and_pdf.py
Genera la Guía Oficial ÚNICA exclusivamente en:
CURSO-ROBOTICA-IA-COMUN

- Cero IDs fósiles o antiguos (únicamente las 2 clases unificadas vigentes).
- Generación de DOCX con párrafos estándar (sin estilos de lista que provoquen
  avisos de recuperación en Microsoft Word).
- Generación paralela de PDF de alta legibilidad con ReportLab.
- Limpieza total de copias redundantes en CURSO-IA.
"""

import os
import sys
import glob
import subprocess

COMUN_DIR = "/Users/externo/Library/Mobile Documents/com~apple~CloudDocs/PERSONAL/CLASES DE TECNOLOGÍA/CURSO-ROBOTICA-IA-COMUN"
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
IA_DIR = os.path.abspath(os.path.join(SCRIPT_DIR, ".."))
PARENT_DIR = os.path.abspath(os.path.join(IA_DIR, ".."))

try:
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "-q"])
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

def limpiar_duplicados_redundantes():
    """Elimina las guías redundantes de CURSO-IA y carpetas satélite."""
    patrones_a_borrar = [
        os.path.join(IA_DIR, "0.GUIA_PUBLICACION_GITHUB_CLASSROOM.*"),
    ]
    for ruta in glob.glob(os.path.join(PARENT_DIR, "*ROBOTICA*")):
        if os.path.abspath(ruta) != os.path.abspath(COMUN_DIR):
            patrones_a_borrar.append(os.path.join(ruta, "0.GUIA_PUBLICACION_GITHUB_CLASSROOM.*"))

    for patron in patrones_a_borrar:
        for f in glob.glob(patron):
            try:
                os.remove(f)
                print(f"🗑️ Eliminada guía redundante: {f}")
            except Exception as e:
                pass

def create_clean_docx(output_path):
    """
    Crea un .docx 100% libre de errores de recuperación.
    Usa párrafos estándar con indentación manual (sin 'List Bullet' ni estilos corruptibles).
    """
    doc = docx.Document()

    # Márgenes de página (2.5 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # TÍTULO
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("GUÍA DE PUBLICACIÓN EN GITHUB Y CLASSROOM")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(17)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(2, 132, 199) # Azul corporativo

    # SUBTÍTULO
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(16)
    r_sub = p_sub.add_run("Workflow de Trabajo en Equipo: Tú (Profesor) y Antigravity (Asistente IA)")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(10)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(100, 116, 139)

    # INTRODUCCIÓN
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_after = Pt(12)
    p_intro.paragraph_format.left_indent = Inches(0.2)
    r_intro_bold = p_intro.add_run("💡 Introducción al Sistema Unificado:\n")
    r_intro_bold.bold = True
    r_intro_bold.font.name = "Arial"
    r_intro_bold.font.size = Pt(10)
    r_intro_bold.font.color.rgb = RGBColor(3, 105, 161)
    r_intro_text = p_intro.add_run(
        "Este flujo de trabajo garantiza que todo el material educativo (teoría, simuladores interactivos, "
        "fichas didácticas y monográficos) viaje desde tu Mac hasta Google Classroom sin gastar espacio en tu Google Drive "
        "y con una experiencia 100% limpia para los alumnos mediante Mapas Interactivos HTML únicos."
    )
    r_intro_text.font.name = "Arial"
    r_intro_text.font.size = Pt(9.5)

    def add_h(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = RGBColor(15, 23, 42)
        return p

    def add_item(text, num="•"):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        r_num = p.add_run(f"{num} ")
        r_num.bold = True
        r_num.font.name = "Arial"
        r_num.font.size = Pt(9.5)
        r_num.font.color.rgb = RGBColor(2, 132, 199)
        r_text = p.add_run(text)
        r_text.font.name = "Arial"
        r_text.font.size = Pt(9.5)
        return p

    # FASE 1
    add_h("FASE 1: PREPARACIÓN LOCAL EN EL MAC 🧑🏻‍🏫 TAREA TUYA")
    add_item("Escribe tu teoría, genera tus PDFs, graba los vídeos y organiza todo en la carpeta correspondiente de tu Mac.", "1.")
    add_item("Nombra los archivos correctamente (ej: T0.DEBATE.pdf, T2.VIDEO.mp4, etc.).", "2.")

    # FASE 2
    add_h("FASE 2: SINCRONIZACIÓN AL SERVIDOR (GITHUB) 🤖 TAREA DE LA IA")
    p_f2 = doc.add_paragraph()
    p_f2.paragraph_format.space_after = Pt(3)
    r_f2_1 = p_f2.add_run("Tú me dices: ")
    r_f2_1.font.name = "Arial"
    r_f2_1.font.size = Pt(9.5)
    r_f2_2 = p_f2.add_run('"Ya he metido todos los archivos nuevos en la carpeta del Mac, súbelos a GitHub".')
    r_f2_2.italic = True
    r_f2_2.font.name = "Arial"
    r_f2_2.font.size = Pt(9.5)
    add_item("Añadir los archivos creados o modificados y hacer el commit.", "•")
    add_item("Subirlos a GitHub (git push origin main) en un único comando limpio.", "•")
    add_item("Objetivo: Alojarlos gratis en GitHub Pages para enlazarlos desde Classroom sin gastar espacio en tu Google Drive. Al actualizar un archivo, solo se vuelve a subir a GitHub y se actualiza al instante para todos los alumnos.", "•")

    # FASE 3
    add_h("FASE 3: CREACIÓN DEL PANEL UNIFICADO DE ENLACES 🤖 TAREA DE LA IA")
    add_item("Yo (la IA) genero el archivo maestro CURSO_IA_UNIFICADO.csv (en PANELES_CSV/).", "•")
    add_item("Arquitectura de Mapa Único: En lugar de saturar Classroom con docenas de enlaces sueltos, cada sección cuenta con un Mapa Interactivo HTML único (MAPA_MONOGRAFICOS.html y MAPA_SESION_XX.html) que agrupa fichas, simuladores y recursos interactivos sin rehacer publicaciones en Classroom.", "•")

    # FASE 4
    add_h("FASE 4: PREPARACIÓN DE GOOGLE CLASSROOM 🧑🏻‍🏫 TAREA TUYA")
    add_item("Entra en classroom.google.com con tu cuenta de profesor.", "1.")
    add_item("Trabajamos exclusivamente con las dos Clases Maestras Unificadas Anuales:", "2.")

    # TABLA EXCLUSIVA DE LAS 2 CLASES VIGENTES
    t = doc.add_table(rows=3, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["CURSO", "NOMBRE EN CLASSROOM", "ID REAL DE LA CLASE"]
    for c_idx, h in enumerate(headers):
        cell = t.cell(0, c_idx)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(2, 132, 199)

    clases = [
        ("INTELIGENCIA ARTIFICIAL", "IA.", "890436514251"),
        ("ROBÓTICA", "ROBÓTICA.", "890435922022")
    ]
    for r_idx, (c, n, id_r) in enumerate(clases, start=1):
        for c_idx, val in enumerate([c, n, id_r]):
            cell = t.cell(r_idx, c_idx)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            if c_idx == 2:
                r.bold = True
                r.font.color.rgb = RGBColor(16, 185, 129)

    # FASE 5
    add_h("FASE 5: INGESTA DE DATOS (GOOGLE SHEETS) 🧑🏻‍🏫 TAREA TUYA")
    add_item("Abre tu navegador en Modo Incógnito / Ventana Privada (vital para evitar conflictos entre cuentas de Google).", "1.")
    add_item("Entra a Google Drive y abre tu hoja de cálculo permanente: 0.PANEL_CLASSROOM.", "2.")
    add_item("Ve a Archivo > Importar > Subir y arrastra PANELES_CSV/CURSO_IA_UNIFICADO.csv.", "3.")
    add_item("Selecciona 'Reemplazar hoja de cálculo actual' y pulsa Importar.", "4.")
    add_item("En la Columna A (ID_CURSO), pega el ID real: 890436514251 y arrastra hacia abajo para que todas las filas tengan el ID.", "5.")

    # FASE 6
    add_h("FASE 6: PUBLICACIÓN AUTOMÁTICA 🧑🏻‍🏫 TAREA TUYA")
    add_item("Dentro de la hoja de Google Sheets, ve a Extensiones > Apps Script.", "1.")
    add_item("Selecciona la función publicarTodoConTemas y pulsa Ejecutar.", "2.")
    add_item("El script creará automáticamente el tema MONOGRAFICOS y los temas SESIÓN 1 a SESIÓN 60 con sus respectivos mapas interactivos en modo BORRADOR (DRAFT).", "3.")
    add_item("¡Listo! Ya puedes ir publicando cada sesión a los alumnos conforme avance el calendario escolar.", "4.")

    # ANEXO 1
    add_h("🚨 ANEXO 1: EL BOTÓN DEL PÁNICO (RESET COMPLETO DE CLASE) 🧑🏻‍🏫 TAREA TUYA")
    p_pan = doc.add_paragraph()
    p_pan.paragraph_format.space_before = Pt(3)
    p_pan.paragraph_format.space_after = Pt(3)
    r_pan_title = p_pan.add_run("Si la subida ha quedado desfasada o necesitas borrar TODO de Classroom para empezar de cero:\n")
    r_pan_title.bold = True
    r_pan_title.font.name = "Arial"
    r_pan_title.font.size = Pt(9.5)
    r_pan_title.font.color.rgb = RGBColor(220, 38, 38)
    add_item("Asegúrate de que el ID del curso está puesto en la celda A2 de la hoja: 890436514251.", "1.")
    add_item("Ve a Extensiones > Apps Script, selecciona la función borrarAbsolutamenteTodo y pulsa Ejecutar.", "2.")
    add_item("Toda la clase (Temas, Materiales y Borradores) quedará completamente vacía en segundos.", "3.")

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"✅ Microsoft Word DOCX limpio generado en: {output_path}")

def create_clean_pdf(output_path):
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, HRFlowable
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        rightMargin=40, leftMargin=40, topMargin=30, bottomMargin=30
    )
    styles = getSampleStyleSheet()

    style_title = ParagraphStyle(
        'DocTitle', parent=styles['Heading1'],
        fontName='Helvetica-Bold', fontSize=18, leading=22,
        textColor=colors.HexColor('#0284C7'), alignment=1, spaceAfter=4
    )
    style_sub = ParagraphStyle(
        'DocSub', parent=styles['Normal'],
        fontName='Helvetica-Oblique', fontSize=10, leading=13,
        textColor=colors.HexColor('#64748B'), alignment=1, spaceAfter=14
    )
    style_h2 = ParagraphStyle(
        'DocH2', parent=styles['Heading2'],
        fontName='Helvetica-Bold', fontSize=11.5, leading=14.5,
        textColor=colors.HexColor('#0F172A'), spaceBefore=11, spaceAfter=3
    )
    style_body = ParagraphStyle(
        'DocBody', parent=styles['Normal'],
        fontName='Helvetica', fontSize=8.5, leading=11.5,
        textColor=colors.HexColor('#1E293B'), spaceAfter=3
    )
    style_alert = ParagraphStyle(
        'DocAlert', parent=styles['Normal'],
        fontName='Helvetica', fontSize=8.5, leading=11.5,
        textColor=colors.HexColor('#991B1B')
    )

    story = [
        Paragraph("GUÍA DE PUBLICACIÓN EN GITHUB Y CLASSROOM", style_title),
        Paragraph("Workflow de Trabajo en Equipo: Tú (Profesor) y Antigravity (Asistente IA)", style_sub),
        HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#0284C7'), spaceAfter=8),
        
        Table([[Paragraph("<b>💡 INTRODUCCIÓN AL SISTEMA UNIFICADO:</b><br/>Este flujo de trabajo garantiza que todo el material educativo viaje desde tu Mac hasta Google Classroom sin coste de almacenamiento en Google Drive mediante <b>Mapas Interactivos HTML únicos</b>.", style_body)]],
              colWidths=[530], style=[
                  ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F0F9FF')),
                  ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#BAE6FD')),
                  ('TOPPADDING', (0,0), (-1,-1), 6),
                  ('BOTTOMPADDING', (0,0), (-1,-1), 6),
                  ('LEFTPADDING', (0,0), (-1,-1), 10),
                  ('RIGHTPADDING', (0,0), (-1,-1), 10),
              ]),
        Spacer(1, 6),

        Paragraph("FASE 1: PREPARACIÓN LOCAL EN EL MAC 🧑🏻‍🏫 TAREA TUYA", style_h2),
        Paragraph("1. Escribe tu teoría, genera tus PDFs, graba los vídeos y organiza todo en la carpeta correspondiente de tu Mac.", style_body),
        Paragraph("2. Nombra los archivos correctamente (ej: <code>T0.DEBATE.pdf</code>, <code>T2.VIDEO.mp4</code>, etc.).", style_body),

        Paragraph("FASE 2: SINCRONIZACIÓN AL SERVIDOR (GITHUB) 🤖 TAREA DE LA IA", style_h2),
        Paragraph("Tú me dices: <i>'Ya he metido todos los archivos nuevos en la carpeta del Mac, súbelos a GitHub'</i>.", style_body),
        Paragraph("Yo (la IA) me encargo de añadir los archivos, hacer el commit y subirlos a la nube (<code>git push</code>) en un único comando limpio.", style_body),
        Paragraph("<b>Objetivo:</b> Alojarlos gratis en GitHub Pages para enlazarlos desde Classroom sin gastar espacio en tu Google Drive.", style_body),

        Paragraph("FASE 3: CREACIÓN DEL PANEL UNIFICADO DE ENLACES 🤖 TAREA DE LA IA", style_h2),
        Paragraph("Yo (la IA) escaneo las carpetas y genero el archivo maestro <b><code>CURSO_IA_UNIFICADO.csv</code></b> (en <code>PANELES_CSV/</code>).", style_body),
        Paragraph("<b>Arquitectura de Mapa Único:</b> En lugar de saturar Classroom con docenas de enlaces sueltos, cada sección cuenta con un <b>Mapa Interactivo HTML único</b> (<code>MAPA_MONOGRAFICOS.html</code> y <code>MAPA_SESION_XX.html</code>).", style_body),

        Paragraph("FASE 4: PREPARACIÓN DE GOOGLE CLASSROOM 🧑🏻‍🏫 TAREA TUYA", style_h2),
        Paragraph("Entra en <code>classroom.google.com</code> con tu cuenta de profesor. Trabajamos exclusivamente con las <b>dos Clases Maestras Unificadas Anuales</b>:", style_body),
        Spacer(1, 3),
        Table([
            [Paragraph("<b>CURSO</b>", style_body), Paragraph("<b>NOMBRE EN CLASSROOM</b>", style_body), Paragraph("<b>ID REAL DE LA CLASE</b>", style_body)],
            [Paragraph("INTELIGENCIA ARTIFICIAL", style_body), Paragraph("<b>IA.</b>", style_body), Paragraph("<b>890436514251</b>", style_body)],
            [Paragraph("ROBÓTICA", style_body), Paragraph("<b>ROBÓTICA.</b>", style_body), Paragraph("<b>890435922022</b>", style_body)],
        ], colWidths=[180, 170, 180], style=[
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0284C7')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
            ('BACKGROUND', (0,1), (-1,1), colors.HexColor('#F8FAFC')),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ]),
        Spacer(1, 6),

        Paragraph("FASE 5: INGESTA DE DATOS (GOOGLE SHEETS) 🧑🏻‍🏫 TAREA TUYA", style_h2),
        Paragraph("1. Abre tu navegador en <b>Modo Incógnito / Ventana Privada</b> (evita conflictos entre cuentas de Google).", style_body),
        Paragraph("2. Entra a Google Drive y abre tu hoja permanente: <b><code>0.PANEL_CLASSROOM</code></b>.", style_body),
        Paragraph("3. Ve a <b>Archivo > Importar > Subir</b> y arrastra <b><code>PANELES_CSV/CURSO_IA_UNIFICADO.csv</code></b>.", style_body),
        Paragraph("4. Selecciona <b>Reemplazar hoja de cálculo actual</b> y pulsa Importar.", style_body),
        Paragraph("5. En la <b>Columna A (<code>ID_CURSO</code>)</b>, pega el ID <b><code>890436514251</code></b> y arrástralo hacia abajo.", style_body),

        Paragraph("FASE 6: PUBLICACIÓN AUTOMÁTICA 🧑🏻‍🏫 TAREA TUYA", style_h2),
        Paragraph("1. En la hoja de Google Sheets, ve a <b>Extensiones > Apps Script</b>.", style_body),
        Paragraph("2. Selecciona la función <b><code>publicarTodoConTemas</code></b> y pulsa <b>Ejecutar</b>.", style_body),
        Paragraph("3. El script creará automáticamente el tema <code>MONOGRAFICOS</code> y los temas <code>SESIÓN 1</code> a <code>SESIÓN 60</code> con sus mapas interactivos en modo BORRADOR (DRAFT).", style_body),

        Spacer(1, 6),
        Table([[Paragraph("<b>🚨 ANEXO 1: EL BOTÓN DEL PÁNICO (RESET COMPLETO DE CLASE)</b><br/>"
                          "Si necesitas borrar TODO de Classroom para empezar de cero:<br/>"
                          "1. Asegúrate de que el ID está puesto en la celda <b>A2</b> de la hoja: <b>890436514251</b>.<br/>"
                          "2. En Apps Script, selecciona la función <b><code>borrarAbsolutamenteTodo</code></b> y pulsa <b>Ejecutar</b>.<br/>"
                          "3. Toda la clase quedará completamente limpia en segundos.", style_alert)]],
              colWidths=[530], style=[
                  ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#FEF2F2')),
                  ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#FCA5A5')),
                  ('TOPPADDING', (0,0), (-1,-1), 6),
                  ('BOTTOMPADDING', (0,0), (-1,-1), 6),
                  ('LEFTPADDING', (0,0), (-1,-1), 10),
                  ('RIGHTPADDING', (0,0), (-1,-1), 10),
              ]),
    ]
    doc.build(story)
    print(f"✅ Guía en PDF limpia generada en: {output_path}")

def main():
    limpiar_duplicados_redundantes()

    dest_docx_comun = os.path.join(COMUN_DIR, "0.GUIA_PUBLICACION_GITHUB_CLASSROOM.docx")
    dest_pdf_comun = os.path.join(COMUN_DIR, "0.GUIA_PUBLICACION_GITHUB_CLASSROOM.pdf")

    create_clean_docx(dest_docx_comun)
    create_clean_pdf(dest_pdf_comun)

if __name__ == "__main__":
    main()
